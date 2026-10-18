"""Tests for Rapikan Auto-Fixer Engine.

Non-negotiable Invariants:
1. File asli tidak pernah ditimpa (salinan baru dibuat di out_path).
2. Perubahan style menggunakan style Word asli (pStyle), bukan format langsung per run.
3. Teks naskah identik sebelum/sesudah fix. Aturan changes_text=True (B002) DILARANG
   berjalan di mode batch 'Rapikan Semua Sekaligus'. Hanya usulan yang disetujui satu per satu.
   Normalisasi spasi ganda (W010) menjaga hash teks setelah normalisasi whitespace identik.
4. Jika analisis ulang masih menemukan temuan untuk aturan yang diperbaiki, atau muncul
   temuan baru, lakukan ROLLBACK otomatis: hapus out_path dan batalkan status fixed.
"""

from pathlib import Path
import pytest
from docx import Document
from docx.shared import Pt
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from rapikan.converter.docx_converter import convert_docx
from rapikan.fixer.orchestrator import DocumentFixer, RollbackError
from rapikan.rules.campus_profile import DEFAULT_PROFILE
from rapikan.rules.word.W001_fake_heading import W001FakeHeading
from rapikan.rules.word.W006_table_no_header import W006TableNoHeader
from rapikan.rules.word.W008_manual_spacing import W008ManualSpacing
from rapikan.rules.word.W010_double_space import W010DoubleSpace
from rapikan.rules.word.B002_common_typos import B002CommonTypos


def test_w001_auto_fix_applies_genuine_heading_style(tmp_path: Path):
    """W001 fix: Replaces Normal style on bold+large text with proper Heading 1 style."""
    src_docx = tmp_path / "original_w001.docx"
    fixed_docx = tmp_path / "fixed_w001.docx"

    doc = Document()
    p = doc.add_paragraph("BAB I PENDAHULUAN")
    p.runs[0].bold = True
    p.runs[0].font.size = Pt(14)
    doc.save(str(src_docx))

    ir_before = convert_docx(src_docx)
    rule = W001FakeHeading()
    findings_before = rule.check(ir_before, DEFAULT_PROFILE)
    assert len(findings_before) == 1
    assert findings_before[0].rule_id == "W001"

    fixer = DocumentFixer()
    summary = fixer.apply_fixes(
        src_path=src_docx,
        out_path=fixed_docx,
        findings=findings_before,
    )

    # Invarian 1: File asli tidak ditimpa & salinan baru terbentuk
    assert src_docx.exists()
    assert fixed_docx.exists()
    assert summary["fixes_applied"] >= 1
    assert summary["reanalysis_clean"] is True

    # Invarian 4: Analisis ulang file hasil harus 0 temuan untuk W001
    ir_after = convert_docx(fixed_docx)
    findings_after = rule.check(ir_after, DEFAULT_PROFILE)
    assert len(findings_after) == 0

    # Invarian 2: Style yang diterapkan adalah Heading resmi, format langsung run dibersihkan
    doc_after = Document(str(fixed_docx))
    assert doc_after.paragraphs[0].style.name in ("Heading 1", "Heading1", "Judul 1")
    assert doc_after.paragraphs[0].runs[0].bold is None


def test_w006_auto_fix_formats_table_header(tmp_path: Path):
    """W006 fix: Formats row 0 of table as a bold, repeating header row."""
    src_docx = tmp_path / "original_w006.docx"
    fixed_docx = tmp_path / "fixed_w006.docx"

    doc = Document()
    tbl = doc.add_table(rows=2, cols=2)
    tbl.rows[0].cells[0].text = "Header 1"
    tbl.rows[0].cells[1].text = "Header 2"
    tbl.rows[1].cells[0].text = "Data 1"
    tbl.rows[1].cells[1].text = "Data 2"
    doc.save(str(src_docx))

    ir_before = convert_docx(src_docx)
    rule = W006TableNoHeader()
    findings_before = rule.check(ir_before, DEFAULT_PROFILE)
    assert len(findings_before) == 1

    fixer = DocumentFixer()
    summary = fixer.apply_fixes(
        src_path=src_docx,
        out_path=fixed_docx,
        findings=findings_before,
    )

    assert summary["fixes_applied"] >= 1
    ir_after = convert_docx(fixed_docx)
    findings_after = rule.check(ir_after, DEFAULT_PROFILE)
    assert len(findings_after) == 0


def test_w008_auto_fix_removes_consecutive_blank_paragraphs(tmp_path: Path):
    """W008 fix: Removes consecutive empty paragraphs used for manual spacing."""
    src_docx = tmp_path / "original_w008.docx"
    fixed_docx = tmp_path / "fixed_w008.docx"

    doc = Document()
    doc.add_paragraph("Paragraf pertama.")
    doc.add_paragraph("")
    doc.add_paragraph("   ")
    doc.add_paragraph("Paragraf kedua.")
    doc.save(str(src_docx))

    ir_before = convert_docx(src_docx)
    rule = W008ManualSpacing()
    findings_before = rule.check(ir_before, DEFAULT_PROFILE)
    assert len(findings_before) >= 1

    fixer = DocumentFixer()
    summary = fixer.apply_fixes(
        src_path=src_docx,
        out_path=fixed_docx,
        findings=findings_before,
    )

    assert summary["fixes_applied"] >= 1
    ir_after = convert_docx(fixed_docx)
    findings_after = rule.check(ir_after, DEFAULT_PROFILE)
    assert len(findings_after) == 0


def test_w008_negative_structural_guards_never_deleted(tmp_path: Path):
    """W008 Negative Test: Paragraphs with drawings, math, or separating 2 tables are NEVER deleted."""
    src_docx = tmp_path / "negative_w008.docx"

    doc = Document()
    # 1. Paragraf dengan drawing
    p_drawing = doc.add_paragraph("")
    drawing_xml = parse_xml(f'<w:drawing {nsdecls("w")}><a:graphic {nsdecls("a")}/></w:drawing>')
    p_drawing._p.append(drawing_xml)

    # 2. Dua tabel berurutan dipisahkan satu paragraf kosong
    t1 = doc.add_table(rows=1, cols=1)
    t1.rows[0].cells[0].text = "Tabel 1"
    
    p_separator = doc.add_paragraph("")  # Pemisah tabel wajib bertahan
    
    t2 = doc.add_table(rows=1, cols=1)
    t2.rows[0].cells[0].text = "Tabel 2"

    doc.save(str(src_docx))

    # Analisis IR
    ir = convert_docx(src_docx)
    rule = W008ManualSpacing()
    findings = rule.check(ir, DEFAULT_PROFILE)

    # Nol temuan W008 karena elemen struktural dilindungi
    assert len(findings) == 0


def test_w010_auto_fix_cleans_double_spaces_with_text_hash_invariant(tmp_path: Path):
    """W010 fix: Replaces double spaces while keeping whitespace-normalized hash identical."""
    src_docx = tmp_path / "original_w010.docx"
    fixed_docx = tmp_path / "fixed_w010.docx"

    doc = Document()
    doc.add_paragraph("Ini adalah  kalimat dengan  spasi ganda.")
    doc.save(str(src_docx))

    ir_before = convert_docx(src_docx)
    rule = W010DoubleSpace()
    findings_before = rule.check(ir_before, DEFAULT_PROFILE)
    assert len(findings_before) == 1

    fixer = DocumentFixer()
    summary = fixer.apply_fixes(
        src_path=src_docx,
        out_path=fixed_docx,
        findings=findings_before,
    )

    assert summary["fixes_applied"] >= 1
    ir_after = convert_docx(fixed_docx)
    findings_after = rule.check(ir_after, DEFAULT_PROFILE)
    assert len(findings_after) == 0

    doc_after = Document(str(fixed_docx))
    assert "  " not in doc_after.paragraphs[0].text
    assert doc_after.paragraphs[0].text == "Ini adalah kalimat dengan spasi ganda."


def test_invariant_3_batch_fix_strictly_excludes_text_altering_rules(tmp_path: Path):
    """Invariant 3: B002 (changes_text=True) is STRICTLY EXCLUDED from batch mode."""
    src_docx = tmp_path / "batch_invariant3.docx"
    fixed_docx = tmp_path / "fixed_batch_invariant3.docx"

    doc = Document()
    p1 = doc.add_paragraph("BAB I PENDAHULUAN")
    p1.runs[0].bold = True
    p1.runs[0].font.size = Pt(14)
    doc.add_paragraph("Praktek analisa data dilakukan.")
    doc.save(str(src_docx))

    ir_before = convert_docx(src_docx)
    w001_findings = W001FakeHeading().check(ir_before, DEFAULT_PROFILE)
    b002_findings = B002CommonTypos().check(ir_before, DEFAULT_PROFILE)

    assert len(w001_findings) == 1
    assert len(b002_findings) >= 2
    assert b002_findings[0].changes_text is True

    # Jalankan batch fix (mode batch menggabungkan temuan tanpa allow_text_changes)
    fixer = DocumentFixer()
    all_findings = w001_findings + b002_findings
    summary = fixer.apply_fixes(
        src_path=src_docx,
        out_path=fixed_docx,
        findings=all_findings,
        allow_text_changes=False,  # Batch mode!
    )

    # W001 diterapkan, tapi B002 sama sekali TIDAK diterapkan di batch mode
    applied_rules = [d["rule_id"] for d in summary["details"] if d["status"] == "applied"]
    assert "W001" in applied_rules
    assert "B002" not in applied_rules

    # Teks naskah paragraf 2 tetap utuh tidak berubah
    doc_after = Document(str(fixed_docx))
    assert doc_after.paragraphs[1].text == "Praktek analisa data dilakukan."


def test_approved_single_proposal_applies_with_casing(tmp_path: Path):
    """User individually approving a rewrite proposal applies the change with casing preserved."""
    src_docx = tmp_path / "proposal.docx"
    fixed_docx = tmp_path / "proposal_fixed.docx"

    doc = Document()
    doc.add_paragraph("Praktek dan analisa data dilakukan untuk merubah paradigma.")
    doc.save(str(src_docx))

    ir_before = convert_docx(src_docx)
    findings = B002CommonTypos().check(ir_before, DEFAULT_PROFILE)
    assert len(findings) >= 3

    # User menyetujui temuan kata 'Praktek' (index 0) secara spesifik
    fixer = DocumentFixer()
    summary = fixer.apply_fixes(
        src_path=src_docx,
        out_path=fixed_docx,
        findings=[findings[0]],
        allow_text_changes=True,  # Disetujui secara eksplisit per item
    )

    assert summary["fixes_applied"] == 1
    doc_after = Document(str(fixed_docx))
    # Titlecase 'Praktek' -> 'Praktik'
    assert doc_after.paragraphs[0].text.startswith("Praktik dan analisa")


def test_b002_ignores_quotes_urls_and_proper_nouns(tmp_path: Path):
    """B002 never flags words in direct quotes, URLs, bibliography, or mid-sentence capitalized names."""
    src_docx = tmp_path / "b002_exclusions.docx"

    doc = Document()
    # 1. Dalam tanda kutip langsung
    doc.add_paragraph('Budi mengatakan, "praktek itu salah".')
    # 2. URL
    doc.add_paragraph('Kunjungi https://example.com/analisa-data untuk materi.')
    # 3. Nama diri di tengah kalimat (TitleCase bukan awal kalimat)
    doc.add_paragraph('Rumah Sakit Analisa Medika didirikan tahun 1990.')

    doc.save(str(src_docx))

    ir = convert_docx(src_docx)
    rule = B002CommonTypos()
    findings = rule.check(ir, DEFAULT_PROFILE)

    # Nol temuan false positive
    assert len(findings) == 0


def test_idempotency_running_fix_twice_produces_identical_clean_result(tmp_path: Path):
    """Idempotency: Applying fix to an already fixed file produces 0 additional changes."""
    src_docx = tmp_path / "idempotent.docx"
    fix1_docx = tmp_path / "idempotent_v1.docx"
    fix2_docx = tmp_path / "idempotent_v2.docx"

    doc = Document()
    p = doc.add_paragraph("BAB I PENDAHULUAN")
    p.runs[0].bold = True
    p.runs[0].font.size = Pt(14)
    doc.save(str(src_docx))

    fixer = DocumentFixer()
    findings1 = W001FakeHeading().check(convert_docx(src_docx), DEFAULT_PROFILE)
    fixer.apply_fixes(src_docx, fix1_docx, findings1)

    # Run 2: Analisis file hasil
    findings2 = W001FakeHeading().check(convert_docx(fix1_docx), DEFAULT_PROFILE)
    assert len(findings2) == 0

    # Jalankan fix kedua kali
    summary2 = fixer.apply_fixes(fix1_docx, fix2_docx, findings2)
    assert summary2["fixes_applied"] == 0
    assert fix1_docx.stat().st_size == fix2_docx.stat().st_size


def test_fixed_file_reopened_without_error(tmp_path: Path):
    """Fixed document can be cleanly reopened and inspected with python-docx."""
    src_docx = tmp_path / "reopen.docx"
    fixed_docx = tmp_path / "reopen_fixed.docx"

    doc = Document()
    doc.add_paragraph("BAB I PENDAHULUAN").runs[0].bold = True
    doc.save(str(src_docx))

    fixer = DocumentFixer()
    findings = W001FakeHeading().check(convert_docx(src_docx), DEFAULT_PROFILE)
    fixer.apply_fixes(src_docx, fixed_docx, findings)

    # Must reopen cleanly without any XML schema errors
    reopened = Document(str(fixed_docx))
    assert len(reopened.paragraphs) == 1
    assert reopened.paragraphs[0].text == "BAB I PENDAHULUAN"


def test_w001_preserves_meaningful_formatting_italic_subscript_superscript(tmp_path: Path):
    """Requirement 1: W001 cleans fake headings but preserves meaningful run formatting (italic, subscript, superscript)."""
    src_docx = tmp_path / "heading_scientific.docx"
    fixed_docx = tmp_path / "heading_scientific_fixed.docx"

    doc = Document()
    p = doc.add_paragraph()
    r1 = p.add_run("BAB I IDENTIFIKASI ")
    r1.bold = True
    r1.font.size = Pt(14)

    r2 = p.add_run("Streptococcus pneumoniae")  # Meaningful scientific italic
    r2.bold = True
    r2.italic = True
    r2.font.size = Pt(14)

    r3 = p.add_run(" DENGAN METODE PCR")
    r3.bold = True
    r3.font.size = Pt(14)

    doc.save(str(src_docx))

    ir_before = convert_docx(src_docx)
    findings = W001FakeHeading().check(ir_before, DEFAULT_PROFILE)
    assert len(findings) == 1

    fixer = DocumentFixer()
    summary = fixer.apply_fixes(src_docx, fixed_docx, findings)
    assert summary["fixes_applied"] == 1

    # Inspect fixed document
    doc_after = Document(str(fixed_docx))
    p_after = doc_after.paragraphs[0]
    assert p_after.style.name in ("Heading 1", "Heading1", "Judul 1")

    # Verify meaningful formatting was preserved on run 1
    runs = p_after.runs
    assert len(runs) == 3
    assert runs[1].text == "Streptococcus pneumoniae"
    assert runs[1].italic is True  # Meaningful formatting must NOT be stripped!
    assert runs[1].bold is None    # Direct bold stripped in favor of style bold


def test_w010_fixer_preserves_alignment_spaces_before_colon(tmp_path: Path):
    """Requirement 4: W010 fixer never touches alignment spaces before colons (e.g. Nama      : Budi)."""
    src_docx = tmp_path / "alignment.docx"
    fixed_docx = tmp_path / "alignment_fixed.docx"

    doc = Document()
    doc.add_paragraph("Nama      : Budi Santoso")
    doc.save(str(src_docx))

    # Even if finding is manually constructed or triggered, fixer must refuse to modify alignment lines
    from rapikan.ir.finding_schema import Finding, Location, LocationType, Severity, Evidence
    dummy_finding = Finding(
        rule_id="W010",
        rule_name="Spasi Ganda",
        severity=Severity.INFO,
        confidence=1.0,
        location=Location(type=LocationType.PARAGRAPH, id="p_0001", label="Paragraf p_0001"),
        evidence=Evidence(snippet="Nama      : Budi Santoso"),
        auto_fixable=True,
        suggestion="Rapikan spasi",
    )

    fixer = DocumentFixer()
    fixer.apply_fixes(src_docx, fixed_docx, [dummy_finding])

    doc_after = Document(str(fixed_docx))
    assert doc_after.paragraphs[0].text == "Nama      : Budi Santoso"


def test_w008_fixer_strict_whitelist_preserves_tracked_revisions(tmp_path: Path):
    """Requirement 2: Paragraph with tracked revisions (w:ins) is never deleted by W008 fixer."""
    src_docx = tmp_path / "revisions.docx"
    fixed_docx = tmp_path / "revisions_fixed.docx"

    doc = Document()
    p1 = doc.add_paragraph("Paragraf pertama.")
    # Add a paragraph with tracked revision (w:ins) and no text outside revision
    p_rev = doc.add_paragraph()
    ins = parse_xml(r'<w:ins %s w:id="1" w:author="Dosen" w:date="2026-10-18T00:00:00Z"><w:r><w:t>Teks revisi</w:t></w:r></w:ins>' % nsdecls('w'))
    p_rev._p.append(ins)
    p3 = doc.add_paragraph("Paragraf ketiga.")
    doc.save(str(src_docx))

    # Check converter
    ir = convert_docx(src_docx)
    assert ir.has_revisions is True
    assert ir.paragraphs[1].has_revisions is True
    assert ir.paragraphs[1].is_blank_whitelist is False

    # Attempt to run W008 fix
    fixer = DocumentFixer()
    findings = W008ManualSpacing().check(ir, DEFAULT_PROFILE)
    # W008 does not flag p_rev as blank
    assert len(findings) == 0


def test_end_to_end_pipeline_upload_analyze_fix_verify_zero_findings(tmp_path: Path):
    """Full End-to-End Pipeline: Document with mechanical errors -> Fix -> Re-analysis yields ZERO findings."""
    src_docx = tmp_path / "e2e_dirty.docx"
    fixed_docx = tmp_path / "e2e_clean.docx"

    doc = Document()
    # Fake heading
    p_h = doc.add_paragraph("BAB I PENDAHULUAN")
    p_h.runs[0].bold = True
    p_h.runs[0].font.size = Pt(14)

    # Double space in prose
    doc.add_paragraph("Metode penelitian  ini menggunakan  pendekatan kualitatif.")

    # Blank paragraphs
    doc.add_paragraph("")
    doc.add_paragraph("   ")

    # Table missing header
    tbl = doc.add_table(rows=2, cols=2)
    tbl.rows[0].cells[0].text = "No"
    tbl.rows[0].cells[1].text = "Kategori"
    tbl.rows[1].cells[0].text = "1"
    tbl.rows[1].cells[1].text = "Data"

    doc.save(str(src_docx))

    # Step 1: Analyze original
    ir1 = convert_docx(src_docx)
    from rapikan.rules.registry import RuleRegistry
    registry = RuleRegistry()
    registry.load_all()

    rules = [
        W001FakeHeading(),
        W006TableNoHeader(),
        W008ManualSpacing(),
        W010DoubleSpace(),
    ]

    findings1 = []
    for r in rules:
        findings1.extend(r.check(ir1, DEFAULT_PROFILE))
    assert len(findings1) >= 4

    # Step 2: Apply auto-fixes
    fixer = DocumentFixer()
    summary = fixer.apply_fixes(src_docx, fixed_docx, findings1)
    assert summary["fixes_applied"] >= 4
    assert summary["reanalysis_clean"] is True

    # Step 3: Re-analyze fixed document
    ir2 = convert_docx(fixed_docx)
    findings2 = []
    for r in rules:
        findings2.extend(r.check(ir2, DEFAULT_PROFILE))

    # All targeted mechanical errors must now be ZERO!
    assert len(findings2) == 0
